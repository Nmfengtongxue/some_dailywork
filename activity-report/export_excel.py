#!/usr/bin/env python3
"""导出与网页报表一致的 Excel（布局/配色对标 PNG 导出）"""
import io
from datetime import datetime, timedelta, timezone

CN_TZ = timezone(timedelta(hours=8))

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from sync_excel import get_completed

SHEET_NAME = '2026年618活动'

# 与 export-report.js / report.css 一致
HEADER_FILL = PatternFill('solid', fgColor='FF9933')
DEPT_FILL = PatternFill('solid', fgColor='FFFF00')
GRAND_FILL = PatternFill('solid', fgColor='FF9933')
WHITE_FILL = PatternFill('solid', fgColor='FFFFFF')

RATE_FILLS = {
    'employee': {
        0: ('FFCdd2', 'B71C1C'),
        20: ('FFE0B2', 'E65100'),
        50: ('FFF59D', 'F57F17'),
        80: ('C8E6C9', '2E7D32'),
        100: ('66BB6A', 'FFFFFF'),
    },
    'dept': {
        0: ('EF5350', 'FFFFFF'),
        20: ('FF9800', 'FFFFFF'),
        50: ('FBC02D', '333333'),
        80: ('43A047', 'FFFFFF'),
        100: ('2E7D32', 'FFFFFF'),
    },
    'grand': {
        0: ('C62828', 'FFFFFF'),
        20: ('EF6C00', 'FFFFFF'),
        50: ('F9A825', '212121'),
        80: ('388E3C', 'FFFFFF'),
        100: ('1B5E20', 'FFFFFF'),
    },
}

THIN = Side(style='thin', color='CCCCCC')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_BORDER = Border(
    left=Side(style='thin', color='999999'),
    right=Side(style='thin', color='999999'),
    top=Side(style='thin', color='999999'),
    bottom=Side(style='thin', color='999999'),
)


def rate_tier(rate):
    if rate >= 100:
        return 100
    if rate >= 80:
        return 80
    if rate >= 50:
        return 50
    if rate >= 20:
        return 20
    return 0


def calc_rate(completed, target):
    if not target:
        return 0
    return completed / target * 100


def fmt_money(n):
    if n is None or n == 0:
        return ''
    return f'{int(round(n)):,}'


def fmt_rate(rate):
    return f'{rate:.2f}%'


def build_rows(data):
    employees = data.get('employees', [])
    records = data.get('records', [])
    dept_order = ['销售部', '行政部', '医疗中心']
    dept_directors = {'销售部': '陈炳森', '行政部': '范宏云', '医疗中心': '张静'}
    dept_managers = {'销售部': '', '行政部': '范宏云', '医疗中心': '张静'}

    def completed(emp_id):
        emp = next(e for e in employees if e['id'] == emp_id)
        return get_completed(emp, records)

    rows = []
    grand_target = grand_completed = 0

    for dept in dept_order:
        dept_emps = [e for e in employees if e.get('department') == dept]
        if not dept_emps:
            continue
        dept_target = sum(float(e.get('target') or 0) for e in dept_emps)
        dept_done = sum(completed(e['id']) for e in dept_emps)
        rows.append({
            'type': 'dept-summary',
            'name': dept + '合计',
            'target': dept_target,
            'completed': dept_done,
            'rate': calc_rate(dept_done, dept_target),
            'director': dept_directors[dept],
            'manager': dept_managers[dept],
            'department': dept,
        })
        for i, emp in enumerate(dept_emps):
            done = completed(emp['id'])
            rows.append({
                'type': 'employee',
                'seq': emp.get('seq', ''),
                'name': emp['name'],
                'target': float(emp.get('target') or 0),
                'completed': done,
                'rate': calc_rate(done, float(emp.get('target') or 0)),
                'director': dept_directors[dept] if i == 0 else '',
                'manager': dept_managers[dept] if i == 0 else '',
                'department': dept,
                'rowspan': len(dept_emps) if i == 0 else 0,
            })
        grand_target += dept_target
        grand_completed += dept_done

    rows.append({
        'type': 'grand-total',
        'target': grand_target,
        'completed': grand_completed,
        'rate': calc_rate(grand_completed, grand_target),
    })
    return rows


def style_cell(ws, r, c, value, *, fill=None, font=None, align=None, border=None, num_fmt=None):
    cell = ws.cell(r, c, value)
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    if num_fmt:
        cell.number_format = num_fmt
    return cell


def apply_rate_style(ws, r, c, rate, row_type):
    key = 'dept' if row_type == 'dept-summary' else 'grand' if row_type == 'grand-total' else 'employee'
    tier = rate_tier(rate)
    bg, fg = RATE_FILLS[key][tier]
    size = 14 if key == 'grand' and tier != 100 else 15 if key == 'grand' and tier == 100 else 12
    style_cell(
        ws, r, c, fmt_rate(rate),
        fill=PatternFill('solid', fgColor=bg),
        font=Font(name='Microsoft YaHei', size=size, bold=True, color=fg),
        align=Alignment(horizontal='center', vertical='center'),
        border=BORDER,
    )


def export_report_excel(data):
    report_date = (data.get('reportDate') or datetime.now(CN_TZ).strftime('%Y-%m-%d'))[:10]
    rows = build_rows(data)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    # 标题
    ws.merge_cells('A1:G1')
    title = ws['A1']
    title.value = f'统计日期：{report_date}  |  2026年618活动业绩报表'
    title.font = Font(name='Microsoft YaHei', size=11, color='666666')
    title.alignment = Alignment(horizontal='right', vertical='center')

    headers = [
        '序号', '总监', '经理组', '姓名',
        '部门/个人6·18\n已完成业绩总额',
        '6·18任务已完成额度（6月）',
        '任务额完成率%'
    ]
    widths = [6, 12, 10, 14, 22, 22, 14]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        style_cell(
            ws, 2, i, h,
            fill=HEADER_FILL,
            font=Font(name='Microsoft YaHei', size=11, bold=True),
            align=Alignment(horizontal='center', vertical='center', wrap_text=True),
            border=HEADER_BORDER,
        )
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[2].height = 36

    r = 3
    director_merges = []

    for row in rows:
        if row['type'] == 'dept-summary':
            for c in range(1, 8):
                fill = DEPT_FILL
                style_cell(ws, r, c, None, fill=fill, border=BORDER)
            style_cell(ws, r, 2, row['director'], fill=DEPT_FILL, border=BORDER,
                       align=Alignment(horizontal='center', vertical='center'))
            style_cell(ws, r, 3, row['manager'], fill=DEPT_FILL, border=BORDER,
                       align=Alignment(horizontal='center', vertical='center'))
            style_cell(ws, r, 4, row['name'], fill=DEPT_FILL, border=BORDER,
                       font=Font(name='Microsoft YaHei', bold=True),
                       align=Alignment(horizontal='left', vertical='center'))
            style_cell(ws, r, 5, row['target'] or None, fill=DEPT_FILL, border=BORDER,
                       align=Alignment(horizontal='center', vertical='center'))
            style_cell(ws, r, 6, row['completed'] or None, fill=DEPT_FILL, border=BORDER,
                       align=Alignment(horizontal='center', vertical='center'))
            apply_rate_style(ws, r, 7, row['rate'], 'dept-summary')
        elif row['type'] == 'grand-total':
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            for c in range(1, 8):
                style_cell(ws, r, c, None, fill=GRAND_FILL, border=BORDER)
            style_cell(ws, r, 1, '合计', fill=GRAND_FILL, border=BORDER,
                       font=Font(name='Microsoft YaHei', bold=True),
                       align=Alignment(horizontal='center', vertical='center'))
            style_cell(ws, r, 4, '小计', fill=GRAND_FILL, border=BORDER,
                       font=Font(name='Microsoft YaHei', bold=True),
                       align=Alignment(horizontal='center', vertical='center'))
            style_cell(ws, r, 5, row['target'] or None, fill=GRAND_FILL, border=BORDER,
                       font=Font(name='Microsoft YaHei', bold=True),
                       align=Alignment(horizontal='center', vertical='center'))
            style_cell(ws, r, 6, row['completed'] or None, fill=GRAND_FILL, border=BORDER,
                       font=Font(name='Microsoft YaHei', bold=True),
                       align=Alignment(horizontal='center', vertical='center'))
            apply_rate_style(ws, r, 7, row['rate'], 'grand-total')
        else:
            style_cell(ws, r, 1, row['seq'] if row['seq'] != '' else None, fill=WHITE_FILL, border=BORDER,
                       align=Alignment(horizontal='center', vertical='center'))
            if row['rowspan']:
                director_merges.append((r, r + row['rowspan'] - 1, row['director'], row['manager']))
                style_cell(ws, r, 2, row['director'], fill=WHITE_FILL, border=BORDER,
                           align=Alignment(horizontal='center', vertical='center'))
                style_cell(ws, r, 3, row['manager'], fill=WHITE_FILL, border=BORDER,
                           align=Alignment(horizontal='center', vertical='center'))
            style_cell(ws, r, 4, row['name'], fill=WHITE_FILL, border=BORDER,
                       align=Alignment(horizontal='left', vertical='center'))
            style_cell(ws, r, 5, row['target'] or None, fill=WHITE_FILL, border=BORDER,
                       align=Alignment(horizontal='center', vertical='center'))
            style_cell(ws, r, 6, row['completed'] or None, fill=WHITE_FILL, border=BORDER,
                       align=Alignment(horizontal='center', vertical='center'))
            apply_rate_style(ws, r, 7, row['rate'], 'employee')
        ws.row_dimensions[r].height = 22
        r += 1

    for start, end, director, manager in director_merges:
        if end > start:
            ws.merge_cells(start_row=start, start_column=2, end_row=end, end_column=2)
            ws.merge_cells(start_row=start, start_column=3, end_row=end, end_column=3)
        ws.cell(start, 2).value = director
        ws.cell(start, 2).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(start, 3).value = manager
        ws.cell(start, 3).alignment = Alignment(horizontal='center', vertical='center')

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue(), report_date
