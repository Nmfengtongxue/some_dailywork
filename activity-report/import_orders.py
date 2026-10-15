#!/usr/bin/env python3
"""从订单 Excel 导入业绩：清零后逐笔写入填报记录（金额=商品价格×数量，时间取自订单编号）"""
import json
import re
import os
from datetime import datetime, timezone, timedelta
from copy import deepcopy

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_EXCEL = os.path.join(BASE_DIR, '最新数据工作簿2.xlsx')
LATEST_JSON = os.path.join(BASE_DIR, 'data', 'latest.json')
EMPLOYEES_JSON = os.path.join(BASE_DIR, 'data', 'employees.json')

CN_TZ = timezone(timedelta(hours=8))
SKIP_USER_TEXT = frozenset({
    '无', '微信用户', '小程序用户', '肥刚', '琪琪', '知书达李', '汤姆',
    '定西公司管理号', '美年大健康体检中心客服部',
})


def order_timestamp(order_id: str) -> str:
    """订单编号 wxo + 毫秒时间戳 → ISO8601（UTC，与浏览器 toISOString 一致）"""
    num = str(order_id).replace('wxo', '').split('-')[0]
    ms = int(num[:13])
    dt = datetime.fromtimestamp(ms / 1000, tz=CN_TZ)
    return dt.astimezone(timezone.utc).strftime('%Y-%m-%dT%H:%M:%S.%f')[:-3] + 'Z'


def extract_employee(name_set, *texts):
    for text in texts:
        if not text:
            continue
        s = str(text).strip()
        if s in SKIP_USER_TEXT:
            continue
        for name in sorted(name_set, key=len, reverse=True):
            if name in s:
                return name
    return None


def calc_amount(row):
    qty = float(row.get('商品数量') or 0)
    price = float(row.get('商品价格') or 0)
    return int(qty * price) if qty * price == int(qty * price) else round(qty * price, 2)


def build_note(row, amount):
    qty = row.get('商品数量')
    price = row.get('商品价格')
    product = (row.get('商品名称') or '').strip()
    oid = str(row.get('订单编号', '')).strip()
    parts = [f'订单 {oid}']
    if qty is not None and price is not None:
        parts.append(f'{price}×{qty}={amount}')
    if product:
        parts.append(product[:40])
    receiver = row.get('收货人')
    if receiver and str(receiver).strip() not in SKIP_USER_TEXT:
        parts.append(f'收货:{receiver}')
    return ' '.join(parts)


def load_orders(excel_path):
    ws = openpyxl.load_workbook(excel_path, data_only=True)['Sheet1']
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    orders = []
    last_main = None
    for r in range(2, ws.max_row + 1):
        row = {headers[i - 1]: ws.cell(r, i).value for i in range(1, len(headers) + 1)}
        oid = row.get('订单编号')
        if row.get('实付金额(元)') == 86263:
            continue
        if not row.get('商品数量') or not row.get('商品价格'):
            continue
        if not oid or str(oid).strip() == '':
            # 无订单号：视为上一笔订单的续行（如张静 1168×2 + 618×2）
            if last_main:
                row = {**row, '_parentOrderId': last_main['订单编号'], '_parentRow': last_main}
            else:
                continue
        else:
            last_main = row
        orders.append(row)
    return orders


def import_orders(excel_path=None, dry_run=False):
    excel_path = excel_path or DEFAULT_EXCEL
    if not os.path.isfile(excel_path):
        raise FileNotFoundError(excel_path)

    latest = json.load(open(LATEST_JSON, encoding='utf-8'))
    name_set = {e['name'] for e in latest['employees']}

    # 清零
    for emp in latest['employees']:
        emp['initialCompleted'] = 0
    latest['records'] = []
    latest['dailySnapshots'] = {}

    orders = load_orders(excel_path)
    records = []
    skipped = []
    by_employee = {}

    for row in orders:
        parent = row.get('_parentRow')
        oid = str(row.get('订单编号') or row.get('_parentOrderId', '')).strip()
        amount = calc_amount(row)
        if amount <= 0:
            skipped.append((oid, '金额无效'))
            continue

        emp_name = extract_employee(
            name_set,
            row.get('推广人'),
            row.get('用户信息'),
            row.get('收货人'),
        )
        if not emp_name and parent:
            emp_name = extract_employee(
                name_set,
                parent.get('推广人'),
                parent.get('用户信息'),
                parent.get('收货人'),
            )
        if not emp_name:
            skipped.append((oid, f"未匹配员工 推广人={row.get('推广人')} 用户={row.get('用户信息')}"))
            continue

        ts = order_timestamp(oid)
        if row.get('_parentOrderId'):
            # 同单续行：时间略晚 1 秒，保持排序
            from datetime import datetime as dt, timezone as tz
            base = dt.fromisoformat(ts.replace('Z', '+00:00'))
            ts = (base + timedelta(seconds=1)).strftime('%Y-%m-%dT%H:%M:%S.%f')[:-3] + 'Z'

        suffix = '_line2' if row.get('_parentOrderId') and not row.get('订单编号') else ''
        rec_id = f"rec_{oid.replace('wxo', 'o').replace('-', '_')}{suffix}"
        records.append({
            'id': rec_id,
            'employeeId': emp_name,
            'amount': amount,
            'note': build_note(row, amount),
            'timestamp': ts,
            'source': 'import',
            'orderId': oid,
        })
        by_employee[emp_name] = by_employee.get(emp_name, 0) + amount

    records.sort(key=lambda r: r['timestamp'])

    grand = sum(by_employee.values())
    if records:
        last = datetime.fromisoformat(records[-1]['timestamp'].replace('Z', '+00:00')).astimezone(CN_TZ)
        report_date = last.strftime('%Y-%m-%d')
    else:
        report_date = datetime.now(CN_TZ).strftime('%Y-%m-%d')

    latest['records'] = records
    latest['reportDate'] = report_date
    latest['syncedFrom'] = os.path.basename(excel_path)
    latest['syncedAt'] = datetime.now(CN_TZ).isoformat(timespec='seconds')
    latest['_snapshotVersion'] = f'import_orders_{report_date}_{len(records)}'

    if dry_run:
        return {
            'dry_run': True,
            'records': len(records),
            'grand_total': grand,
            'skipped': skipped,
            'by_employee': dict(sorted(by_employee.items(), key=lambda x: -x[1])[:10]),
        }

    with open(LATEST_JSON, 'w', encoding='utf-8') as f:
        json.dump(latest, f, ensure_ascii=False, indent=2)

    emp_data = json.load(open(EMPLOYEES_JSON, encoding='utf-8'))
    emp_data['employees'] = deepcopy(latest['employees'])
    with open(EMPLOYEES_JSON, 'w', encoding='utf-8') as f:
        json.dump(emp_data, f, ensure_ascii=False, indent=2)

    backup_path = os.path.join(BASE_DIR, 'data', 'import_result.json')
    with open(backup_path, 'w', encoding='utf-8') as f:
        json.dump({
            'importedAt': latest['syncedAt'],
            'source': excel_path,
            'recordCount': len(records),
            'grandTotal': grand,
            'skipped': skipped,
            'records': records,
        }, f, ensure_ascii=False, indent=2)

    return {
        'records': len(records),
        'grand_total': grand,
        'skipped': skipped,
        'report_date': report_date,
        'version': latest['_snapshotVersion'],
    }


if __name__ == '__main__':
    import sys
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_EXCEL
    result = import_orders(path)
    print(json.dumps(result, ensure_ascii=False, indent=2))
